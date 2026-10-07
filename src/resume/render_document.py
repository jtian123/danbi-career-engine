"""ATS-friendly Word renderer, adapted from James's renderer patterns.

Run with the bundled document runtime. No experience-specific James rules or data.
"""
from __future__ import annotations
import json
import re
import sys
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def anchor_re(anchor):
    return re.compile(r'(?<!\w)' + re.escape(anchor) + r'(?!\w)')


class Renderer:
    def __init__(self, config, links):
        self.cfg = config
        self.links = links
        self.linked = set()

    def run(self, p, text, bold=False, italic=False, size=None):
        r = p.add_run(text)
        r.bold, r.italic = bold, italic
        r.font.name = self.cfg['font']
        r.font.size = Pt(size or self.cfg['font_size'])
        r.font.color.rgb = RGBColor(0, 0, 0)
        return r

    def link(self, p, text, url, bold=False, size=None):
        relation = p.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
        link = OxmlElement('w:hyperlink'); link.set(qn('r:id'), relation)
        run = OxmlElement('w:r'); props = OxmlElement('w:rPr')
        fonts = OxmlElement('w:rFonts'); fonts.set(qn('w:ascii'), self.cfg['font']); fonts.set(qn('w:hAnsi'), self.cfg['font']); props.append(fonts)
        sz = OxmlElement('w:sz'); sz.set(qn('w:val'), str(int((size or self.cfg['font_size']) * 2))); props.append(sz)
        color = OxmlElement('w:color'); color.set(qn('w:val'), '000000'); props.append(color)
        underline = OxmlElement('w:u'); underline.set(qn('w:val'), 'single'); props.append(underline)
        if bold: props.append(OxmlElement('w:b'))
        run.append(props); t = OxmlElement('w:t'); t.text = text; run.append(t); link.append(run); p._p.append(link)

    def emit(self, p, text, bold=False):
        while text:
            matches = [(m.start(), -len(a), a, u, m) for a, u in self.links.items() if u not in self.linked for m in [anchor_re(a).search(text)] if m]
            if not matches:
                self.run(p, text, bold=bold); return
            _, _, anchor, url, match = min(matches)
            if match.start(): self.run(p, text[:match.start()], bold=bold)
            self.link(p, anchor, url, bold=bold); self.linked.add(url); text = text[match.end():]

    def para(self, doc, before=0, after=0, keep=False):
        p = doc.add_paragraph()
        f = p.paragraph_format; f.space_before = Pt(before); f.space_after = Pt(after); f.line_spacing = 1.02
        f.keep_with_next = keep; f.widow_control = True
        f.tab_stops.add_tab_stop(Inches(8.5-2*self.cfg['margin']), WD_TAB_ALIGNMENT.RIGHT)
        return p

    def heading(self, doc, title):
        p = self.para(doc, before=9, after=3, keep=True)
        self.run(p, title.upper(), bold=True, size=self.cfg['heading_size'])

    def render(self, resume, output):
        self.linked.clear()
        doc = Document()
        # Some document runtimes ship a Title style with a colored bottom rule.
        # Remove inherited borders as well as direct ones; typography supplies hierarchy.
        for border in list(doc.styles.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
        section = doc.sections[0]; section.page_width = Inches(8.5); section.page_height = Inches(11)
        for key in ('top_margin','bottom_margin','left_margin','right_margin'):setattr(section,key,Inches(self.cfg['margin']))
        normal = doc.styles['Normal']; normal.font.name = self.cfg['font']; normal.font.size = Pt(self.cfg['font_size']); normal.font.color.rgb = RGBColor(0,0,0)
        doc.core_properties.author = resume['name']; doc.core_properties.title = resume['name']+' - '+resume['headline']; doc.core_properties.subject = 'Resume'; doc.core_properties.last_modified_by = resume['name']
        p = self.para(doc, after=2, keep=True); p.style = doc.styles['Title']; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.run(p,resume['name'],bold=True,size=self.cfg['name_size'])
        p = self.para(doc,after=3,keep=True);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;self.run(p,resume['headline'],size=self.cfg['font_size'])
        p=self.para(doc,after=4,keep=True);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        contact=resume['contact'];parts=[('email',contact.get('email')),('phone',contact.get('phone')),('linkedin',contact.get('linkedin'))]
        parts=[(k,v) for k,v in parts if v]
        for i,(k,v) in enumerate(parts):
            if i:self.run(p,'  |  ',size=9.5)
            if k=='linkedin':self.link(p,re.sub(r'^https?://(www\.)?','',v).rstrip('/'),v,size=9.5)
            else:self.run(p,v,size=9.5)
        self.run(self.para(doc,after=3),resume['summary'])
        self.heading(doc,'Education')
        for e in resume['education']:
            p=self.para(doc,keep=True);self.run(p,e['school'],bold=True);self.run(p,'\t'+e['dates'])
            self.run(self.para(doc,after=2,keep=bool(e.get('coursework'))),e['degree']+('  |  GPA '+e['gpa'] if e.get('gpa') else ''))
            for line in e.get('coursework', []):
                self.run(self.para(doc,after=2),line)
        self.heading(doc,'Experience')
        for e in resume['experience']:
            p=self.para(doc,before=4,keep=True);self.emit(p,e['company'],bold=True)
            if e['location']:self.run(p,'\t'+e['location'])
            p=self.para(doc,after=2,keep=True);self.run(p,e['title'],italic=True)
            if e['dates']:self.run(p,'\t'+e['dates'],italic=True)
            for text in e['bullets']:
                p=self.para(doc,after=3);p.style=doc.styles['List Bullet'];p.paragraph_format.left_indent=Inches(.16);p.paragraph_format.first_line_indent=Inches(-.12)
                self.emit(p,text)
        self.heading(doc,'Skills')
        for group in resume['skills']:
            p=self.para(doc,after=2);self.run(p,group['group']+': ',bold=True);self.run(p,'; '.join(group['items']))
        output=Path(output);output.parent.mkdir(parents=True,exist_ok=True);doc.save(output)


if __name__ == '__main__':
    r=json.loads(Path(sys.argv[1]).read_text());c=json.loads(Path(sys.argv[2]).read_text())
    Renderer(c,r['product_links']).render(r,Path(sys.argv[3]))
